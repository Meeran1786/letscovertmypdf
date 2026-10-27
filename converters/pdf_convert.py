#!/usr/bin/env python3
"""PDF -> DOCX / XLSX converters used by server.js.

Usage: pdf_convert.py docx|xlsx <input.pdf> <output>
Exits non-zero with a message on stderr when conversion fails.
"""
import re
import sys


def to_docx(src, dst):
    from pdf2docx import Converter
    cv = Converter(src)
    try:
        cv.convert(dst)
    finally:
        cv.close()


NUM_RE = re.compile(r'^-?[\d,]*\.?\d+$')


def cell_value(v):
    if v is None:
        return None
    v = str(v).strip()
    if NUM_RE.match(v) and v.count(',') * 4 <= len(v):
        try:
            n = float(v.replace(',', ''))
            return int(n) if n.is_integer() and '.' not in v else n
        except ValueError:
            pass
    return v


def to_xlsx(src, dst):
    import pdfplumber
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(src) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            # Ruled tables first; fall back to text-aligned columns for unruled layouts.
            tables = page.extract_tables()
            if not tables:
                tables = page.extract_tables({'vertical_strategy': 'text', 'horizontal_strategy': 'text'})
            ws = wb.create_sheet(f'Page {i}')
            row = 1
            for table in tables:
                for r in table:
                    if not any((c or '').strip() for c in r):
                        continue
                    for col, v in enumerate(r, start=1):
                        ws.cell(row=row, column=col, value=cell_value(v))
                    row += 1
                row += 1  # blank line between tables
            if row == 1:
                text = page.extract_text() or ''
                for line in text.splitlines():
                    ws.cell(row=row, column=1, value=line)
                    row += 1
            for column in ws.columns:
                width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
                ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), 60)
    if not wb.worksheets:
        wb.create_sheet('Page 1')
    wb.save(dst)


def main():
    if len(sys.argv) != 4 or sys.argv[1] not in ('docx', 'xlsx'):
        print(__doc__, file=sys.stderr)
        return 2
    mode, src, dst = sys.argv[1:]
    try:
        (to_docx if mode == 'docx' else to_xlsx)(src, dst)
    except Exception as e:  # surface a readable message to the API caller
        print(f'{type(e).__name__}: {e}', file=sys.stderr)
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
